"""
Reporting module for SmartStock AI.
Generates an executive-ready, multi-tab Excel workbook using openpyxl with
custom navy styling, number formatting, auto-adjusted column widths, and frozen header rows.
"""

import io
from typing import Dict, Any, List
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.forecasting import prepare_product_series, forecast_exponential_smoothing, evaluate_models_chronological


def style_excel_sheet(ws, title: str = ""):
    """Applies enterprise navy styling, frozen panes, and borders to an openpyxl worksheet."""
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=9)
    border_side = Side(style="thin", color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Freeze header row
    ws.freeze_panes = "A2"

    max_row = ws.max_row
    max_col = ws.max_column

    # Format Header Row
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws.row_dimensions[1].height = 28

    # Format Data Rows
    for row_idx in range(2, max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = cell_border
            if is_even:
                cell.fill = alt_fill

            # Apply sensible number formatting based on column header
            header_val = str(ws.cell(row=1, column=col_idx).value or "").lower()
            if any(term in header_val for term in ["cost", "value", "price", "amount", "po_val"]):
                cell.number_format = '"₹"#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif any(term in header_val for term in ["qty", "stock", "units", "point", "orders", "days", "count"]):
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif any(term in header_val for term in ["pct", "rate", "wape", "bias", "margin"]):
                cell.number_format = '0.0%'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif any(term in header_val for term in ["date"]):
                cell.number_format = 'yyyy-mm-dd'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif any(term in header_val for term in ["demand", "mae", "score", "dos"]):
                cell.number_format = '0.0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        adjusted_width = min(40, max(max_len + 3, 12))
        ws.column_dimensions[col_letter].width = adjusted_width


def generate_executive_excel_report(
    inventory_df: pd.DataFrame,
    raw_df: pd.DataFrame,
    service_level_str: str = "95% (Z = 1.65)",
    holding_rate: float = 0.20,
) -> io.BytesIO:
    """
    Creates an Excel report with all 7 required sheets:
    1. Inventory Summary
    2. Demand Forecasts
    3. Forecast Accuracy
    4. Reorder Recommendations
    5. Stockout Risk
    6. Excess Inventory
    7. Methodology and Assumptions
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # ---------------------------------------------------------
    # Sheet 1: Inventory Summary
    # ---------------------------------------------------------
    ws_summary = wb.create_sheet(title="Inventory_Summary")
    summary_cols = [
        "Product_ID", "Product_Name", "Category", "On_Hand_Stock", "On_Order_Stock",
        "Backorders", "Inventory_Position", "Avg_Daily_Demand", "Lead_Time_Days",
        "Safety_Stock", "Reorder_Point", "Days_of_Supply", "Unit_Cost",
        "Inventory_Value", "Annual_Holding_Cost", "Inventory_Status"
    ]
    summary_data = inventory_df[summary_cols].copy()
    ws_summary.append(summary_cols)
    for _, row in summary_data.iterrows():
        ws_summary.append(row.tolist())
    style_excel_sheet(ws_summary)

    # ---------------------------------------------------------
    # Sheet 2: Demand Forecasts (14-Day Horizon for all SKUs)
    # ---------------------------------------------------------
    ws_forecast = wb.create_sheet(title="Demand_Forecasts")
    forecast_rows = []
    for sku_id in inventory_df["Product_ID"].unique():
        sku_name = inventory_df.loc[inventory_df["Product_ID"] == sku_id, "Product_Name"].values[0]
        cat = inventory_df.loc[inventory_df["Product_ID"] == sku_id, "Category"].values[0]
        series = prepare_product_series(raw_df, sku_id)
        if len(series) >= 14:
            fc_series, _ = forecast_exponential_smoothing(series, horizon=14)
            for dt, val in fc_series.items():
                forecast_rows.append({
                    "Date": dt.strftime("%Y-%m-%d"),
                    "Product_ID": sku_id,
                    "Product_Name": sku_name,
                    "Category": cat,
                    "Forecast_Units": round(val, 1),
                })
    df_fc = pd.DataFrame(forecast_rows)
    ws_forecast.append(list(df_fc.columns))
    for _, row in df_fc.iterrows():
        ws_forecast.append(row.tolist())
    style_excel_sheet(ws_forecast)

    # ---------------------------------------------------------
    # Sheet 3: Forecast Accuracy (Backtest on last 14 days)
    # ---------------------------------------------------------
    ws_accuracy = wb.create_sheet(title="Forecast_Accuracy")
    acc_rows = []
    for sku_id in inventory_df["Product_ID"].unique():
        sku_name = inventory_df.loc[inventory_df["Product_ID"] == sku_id, "Product_Name"].values[0]
        series = prepare_product_series(raw_df, sku_id)
        eval_res = evaluate_models_chronological(series, test_horizon=14, ma_window=14)
        if eval_res.get("has_data"):
            acc_rows.append({
                "Product_ID": sku_id,
                "Product_Name": sku_name,
                "Model_Tested": eval_res.get("exp_model_name", "Holt-Winters"),
                "MAE_Units": eval_res["exp_metrics"]["MAE"],
                "WAPE_Pct": eval_res["exp_metrics"]["WAPE"] / 100.0,
                "Bias_Pct": eval_res["exp_metrics"]["Bias"] / 100.0,
                "MA_Baseline_WAPE": eval_res["ma_metrics"]["WAPE"] / 100.0,
            })
    df_acc = pd.DataFrame(acc_rows)
    ws_accuracy.append(list(df_acc.columns))
    for _, row in df_acc.iterrows():
        ws_accuracy.append(row.tolist())
    style_excel_sheet(ws_accuracy)

    # ---------------------------------------------------------
    # Sheet 4: Reorder Recommendations
    # ---------------------------------------------------------
    ws_reorder = wb.create_sheet(title="Reorder_Recommendations")
    reorder_subset = inventory_df[inventory_df["Recommended_Order_Qty"] > 0].copy()
    reorder_subset["Est_PO_Value"] = reorder_subset["Recommended_Order_Qty"] * reorder_subset["Unit_Cost"]
    reorder_cols = [
        "Product_ID", "Product_Name", "Category", "On_Hand_Stock", "On_Order_Stock",
        "Inventory_Position", "Reorder_Point", "Recommended_Order_Qty",
        "Unit_Cost", "Est_PO_Value", "Lead_Time_Days", "Recommendation_Reason"
    ]
    df_reorder = reorder_subset[reorder_cols] if not reorder_subset.empty else pd.DataFrame(columns=reorder_cols)
    ws_reorder.append(reorder_cols)
    for _, row in df_reorder.iterrows():
        ws_reorder.append(row.tolist())
    style_excel_sheet(ws_reorder)

    # ---------------------------------------------------------
    # Sheet 5: Stockout Risk
    # ---------------------------------------------------------
    ws_stockout = wb.create_sheet(title="Stockout_Risk")
    stockout_subset = inventory_df[inventory_df["Inventory_Status"] == "Potential Stockout Risk"].copy()
    stockout_cols = [
        "Product_ID", "Product_Name", "Category", "On_Hand_Stock", "Backorders",
        "Avg_Daily_Demand", "Lead_Time_Demand", "Days_of_Supply", "Lead_Time_Days", "Recommendation_Reason"
    ]
    df_stockout = stockout_subset[stockout_cols] if not stockout_subset.empty else pd.DataFrame(columns=stockout_cols)
    ws_stockout.append(stockout_cols)
    for _, row in df_stockout.iterrows():
        ws_stockout.append(row.tolist())
    style_excel_sheet(ws_stockout)

    # ---------------------------------------------------------
    # Sheet 6: Excess Inventory
    # ---------------------------------------------------------
    ws_excess = wb.create_sheet(title="Excess_Inventory")
    excess_subset = inventory_df[inventory_df["Inventory_Status"] == "Excess Inventory"].copy()
    excess_subset["Excess_Units"] = excess_subset["On_Hand_Stock"] - (2.0 * excess_subset["Reorder_Point"])
    excess_subset["Excess_Units"] = excess_subset["Excess_Units"].clip(lower=0)
    excess_subset["Capital_Locked_Up"] = excess_subset["Excess_Units"] * excess_subset["Unit_Cost"]
    excess_cols = [
        "Product_ID", "Product_Name", "Category", "On_Hand_Stock", "Reorder_Point",
        "Excess_Units", "Days_of_Supply", "Unit_Cost", "Capital_Locked_Up", "Recommendation_Reason"
    ]
    df_excess = excess_subset[excess_cols] if not excess_subset.empty else pd.DataFrame(columns=excess_cols)
    ws_excess.append(excess_cols)
    for _, row in df_excess.iterrows():
        ws_excess.append(row.tolist())
    style_excel_sheet(ws_excess)

    # ---------------------------------------------------------
    # Sheet 7: Methodology and Assumptions
    # ---------------------------------------------------------
    ws_method = wb.create_sheet(title="Methodology_Assumptions")
    method_data = [
        ["Parameter / Concept", "Formula / Specification", "Supply Chain Rationale & Assumptions"],
        ["Target Service Level", service_level_str, "Probability of not stocking out during replenishment cycle under normal distribution."],
        ["Inventory Position (IP)", "IP = On-Hand + On-Order - Backorders", "True accounting of net availability to prevent double-ordering pipeline stock."],
        ["Expected Lead-Time Demand (LTD)", "LTD = Average Daily Demand * Lead Time (days)", "Mean volume demanded while waiting for vendor shipment arrival."],
        ["Safety Stock (SS)", "SS = Z * sqrt( L * sigma_d^2 + d^2 * sigma_L^2 )", "Protects against both demand variability and supplier delivery delays."],
        ["Reorder Point (ROP)", "ROP = Lead-Time Demand + Safety Stock", "Stock threshold below which a purchase order must be generated immediately."],
        ["Replenishment Quantity", "Order Qty = max(0, ROP + (d * Review_Period) - IP)", "Periodic (R, s, S) base stock policy ensuring sufficient coverage."],
        ["Annual Holding Cost Rate", f"{holding_rate * 100:.1f}% per annum", "Cost of working capital, warehousing storage, insurance, and obsolescence risk."],
        ["Forecast Evaluation Metric", "WAPE = sum(|Actual - Pred|) / sum(Actual)", "Preferred over MAPE in FMCG retail because zero sales days do not cause division-by-zero errors."],
        ["Data Disclaimer", "Synthetic & Simulated Dataset", "Generated for academic demonstration and ITC internship preparation. Does not represent proprietary commercial data."],
    ]
    for row in method_data:
        ws_method.append(row)
    style_excel_sheet(ws_method)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

